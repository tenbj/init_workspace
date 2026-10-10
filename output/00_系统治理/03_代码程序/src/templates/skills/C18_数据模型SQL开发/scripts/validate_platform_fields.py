#!/usr/bin/env python3
"""只读核验平台六列字段Excel；不连接平台或数据库。"""
import argparse
import json
import re
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

HEADERS = ['排序序号', 'Type', '字段原名', '字段名', '字段释义', '主键']
AUDIT = ['first_inserted_dt', 'etl_dt', 'last_updated_dt']


def type_issue(value):
    if not isinstance(value, str) or not value.strip():
        return 'Type缺失'
    if value != value.lower():
        return 'Type必须小写'
    typ = re.sub(r'\s+', '', value)
    if typ in {'boolean', 'tinyint', 'smallint', 'int', 'integer', 'bigint', 'largeint',
               'float', 'double', 'date', 'datev2', 'datetime', 'datetimev2', 'string'}:
        return None
    if re.fullmatch(r'(?:datetime|datetimev2)\([0-6]\)', typ):
        return None
    if re.fullmatch(r'(?:var)?char\([1-9][0-9]*\)', typ):
        return None
    m = re.fullmatch(r'decimal(?:v3|32|64|128|256)?\(([1-9][0-9]*),([0-9]+)\)', typ)
    if m and int(m[2]) <= int(m[1]):
        return None
    return '类型形式未被静态检查覆盖或精度写法不正确，需人工核对Doris语法'


def validate(excel):
    errors = []
    def add(row, message):
        errors.append({'row': row, 'message': message})
    wb = load_workbook(excel, read_only=True, data_only=False)
    try:
        if len(wb.worksheets) != 1:
            add(0, '导入文件必须只有一个工作表')
        sheet = wb.worksheets[0]
        sheet.reset_dimensions()
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        wb.close()
    while rows and not any(v is not None for v in rows[-1]):
        rows.pop()
    if not rows or list(rows[0]) != HEADERS:
        add(1, '必须是精确的六列表头及顺序')
    data = rows[1:]
    if len(data) < 4:
        add(0, '缺少三个审计字段或业务字段')
    names, labels, keys = [], [], []
    for i, raw in enumerate(data, 2):
        if len(raw) != 6:
            add(i, '字段行必须为六列')
        row = list(raw[:6]) + [None] * max(0, 6-len(raw))
        seq, typ, name, label, comment, key = row
        for value in row:
            if isinstance(value, str) and value.startswith('='):
                add(i, '导入字段信息不得使用Excel公式')
            if isinstance(value, str) and value != value.strip():
                add(i, '单元格含首尾空白，需核对')
        expected = 0 if i < 5 else i - 4
        if isinstance(seq, bool) or not isinstance(seq, (int, float)) or seq != expected:
            add(i, f'排序序号应为{expected}')
        if i < 5 and (name != AUDIT[i-2] or re.sub(r'\s+', '', str(typ)) != 'datetime(3)'):
            add(i, '前三行必须依次为三个datetime(3)审计字段')
        issue = type_issue(typ)
        if issue:
            add(i, issue)
        if not isinstance(name, str) or not re.fullmatch(r'[a-z][a-z0-9]*(?:_[a-z0-9]+)*', name):
            add(i, '字段原名须为小写下划线物理列名')
        if not isinstance(label, str) or not re.search(r'[\u4e00-\u9fff]', label):
            add(i, '字段名须填写中文展示名')
        if not isinstance(comment, str) or not comment.strip():
            add(i, '字段释义须填写数据库COMMENT')
        if key not in (None, '', '是'):
            add(i, '主键标记只能为是或空白')
        if key == '是':
            keys.append(name)
        names.append(str(name).strip())
        labels.append(str(label).strip())
    for label, values in [('字段原名', names), ('字段名', labels)]:
        for value, count in Counter(values).items():
            if count > 1:
                add(0, f'{label}重复：{value}')
    return {'status': 'PASS' if not errors else 'FAIL', 'field_count': len(data),
            'primary_keys': keys, 'errors': errors,
            'scope': '六列、审计/业务序号、类型小写及常用形式、名称唯一、非空与主键标记；未验证DDL/SQL一致性、COMMENT语义、Doris版本兼容或页面排序键'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--excel', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    source, output = Path(args.excel).resolve(), Path(args.output).resolve()
    if source == output or output.suffix.lower() != '.json':
        parser.error('报告必须是独立JSON，不得覆盖Excel')
    result = validate(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'status': result['status'], 'field_count': result['field_count'],
                      'error_count': len(result['errors'])}, ensure_ascii=False))
    return 0 if result['status'] == 'PASS' else 1


if __name__ == '__main__':
    raise SystemExit(main())
