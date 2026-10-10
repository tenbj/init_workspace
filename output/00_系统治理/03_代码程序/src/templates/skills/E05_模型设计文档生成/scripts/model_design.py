"""Inspect Excel evidence and render an AI-reviewed model contract. No SQL execution."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'assets' / '模型设计空白模板.xlsx'
CATALOG = '1.1 模型设计目录'
DETAIL = '1.3 模型实体设计模板'
STATES = {'已确认', '设计建议', '待确认', '存在冲突'}
MODES = {'code', 'requirements', 'rewrite'}
FIELD_KEYS = ['category', 'subcategory', 'name', 'type', 'description', 'technical_key', 'business_key']
META_KEYS = ['grain', 'business_keys', 'technical_keys', 'partition', 'incremental_field',
             'update_strategy', 'frequency', 'schedule', 'joins', 'filters', 'deduplication',
             'metric_rules', 'null_rules', 'history_policy']


def dump(path, data):
    path = Path(path)
    if path.exists():
        raise ValueError(f'拒绝覆盖已有文件，请指定新路径或先执行 B02：{path}')
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')


def inspect(source, output):
    """Preserve exact positions, formulas, cached values and merged-cell anchors."""
    source = Path(source)
    wb = load_workbook(source, data_only=False)
    cached = load_workbook(source, data_only=True)
    result = {'source': str(source.resolve()), 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
              'sheets': [], 'note': '含隐藏页；公式未执行。图片需人工/AI视觉核读，不能据单元格空值认定没有证据。'}
    for ws in wb:
        cells = []
        for row in ws:
            for c in row:
                if c.value is None and c.comment is None and c.hyperlink is None:
                    continue
                link = c.hyperlink
                cells.append({'cell': c.coordinate, 'value': c.value, 'type': c.data_type,
                              'cached': cached[ws.title][c.coordinate].value if c.data_type == 'f' else None,
                              'comment': c.comment.text if c.comment else None,
                              'hyperlink': {'target': link.target, 'location': link.location} if link else None})
        result['sheets'].append({'name': ws.title, 'state': ws.sheet_state, 'dimension': ws.calculate_dimension(),
                                'merges': [str(x) for x in ws.merged_cells.ranges],
                                'image_count': len(ws._images),
                                'validations': [{'range': str(x.sqref), 'type': x.type, 'formula1': x.formula1}
                                                for x in ws.data_validations.dataValidation], 'cells': cells})
    dump(output, result)
    return {'sheets': len(result['sheets']), 'output': str(output)}


def text(value):
    if value is None:
        return '待确认'
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def put(ws, row, col, value):
    value = text(value)
    if len(value) > 32767:
        raise ValueError(f'{ws.title}!{row},{col} 超过Excel单元格32767字符；请拆分并引用完整SQL文件')
    cell = ws.cell(row, col)
    cell.value = value
    cell.data_type = 's'  # SQL, external source strings and =AI(...) are literal evidence, never executable formulas.
    cell.alignment = Alignment(vertical='top', wrap_text=True)
    return cell


def link(ws, address, sheet, target='A1'):
    ws[address].hyperlink = Hyperlink(ref=address, location=f"'{sheet.replace(chr(39), chr(39)*2)}'!{target}")
    ws[address].font = Font(name='微软雅黑', color='0563C1', underline='single')


def validate_contract(data):
    errors, warnings = [], []
    def error(message): errors.append(message)
    if data.get('mode') not in MODES:
        error('mode必须是code / requirements / rewrite')
    if not data.get('title'):
        error('缺少title')
    models = data.get('models', [])
    if not isinstance(models, list) or not models:
        return errors + ['models必须是非空数组'], warnings
    evidence = data.get('evidence', [])
    if not isinstance(evidence, list):
        return errors + ['evidence必须是数组'], warnings
    ids = [e.get('id') for e in evidence]
    if len(ids) != len(set(ids)) or any(not x for x in ids):
        error('evidence.id必须非空且唯一')
    for e in evidence:
        if e.get('state') not in STATES or not e.get('source') or not e.get('locator') or not e.get('claim'):
            error(f'证据{e.get("id")}缺少source/locator/claim或状态非法')
    known = set(ids)
    def check_refs(item, label):
        refs = item.get('evidence', [])
        if not isinstance(refs, list) or not refs or not set(refs) <= known:
            error(f'{label}必须引用已定义的evidence.id')
        if item.get('state') not in STATES:
            error(f'{label}必须显式标记state')
        elif item['state'] != '已确认':
            warnings.append(f'{label}：{item["state"]}')
    tables = []
    for i, model in enumerate(models, 1):
        label = f'模型{i}'
        for k in ['name', 'table', 'layer', 'domain', 'fields', 'design', 'state', 'evidence']:
            if k not in model: error(f'{label}缺少{k}')
        check_refs(model, label)
        tables.append(model.get('table'))
        fields = model.get('fields', [])
        if not fields:
            error(f'{label}无字段')
        names = [f.get('name') for f in fields if f.get('name')]
        if len(names) != len(set(names)):
            error(f'{label}字段名重复；多来源应放进同一字段mappings')
        for j, field in enumerate(fields, 1):
            fl = f'{label}/字段{j}'
            for key in FIELD_KEYS:
                if key not in field: error(f'{fl}缺少{key}')
            if not field.get('description'): error(f'{fl}缺少字段描述')
            check_refs(field, fl)
            for key in ['technical_key', 'business_key']:
                if field.get(key) not in ('是', '否', '待确认', '不适用'):
                    error(f'{fl}/{key}须用是、否、待确认、不适用')
            if (not field.get('name') or not field.get('type')) and field.get('state') == '已确认':
                error(f'{fl}未命名/未知类型不能标为已确认')
            mappings = field.get('mappings', [])
            if not mappings:
                error(f'{fl}缺少mappings；未知来源也须明确记录')
            for mapping in mappings:
                for key in ['branch', 'source_table', 'source_field', 'expression', 'condition']:
                    if key not in mapping: error(f'{fl}/mapping缺少{key}')
                check_refs(mapping, f'{fl}/mapping')
        design = model.get('design', {})
        for key in META_KEYS:
            if key not in design:
                error(f'{label}/design缺少{key}，未知填null，不适用写明原因')
            elif design[key] is None:
                warnings.append(f'{label}/{key}待确认')
        for dk, fk in [('technical_keys', 'technical_key'), ('business_keys', 'business_key')]:
            expected = design.get(dk)
            if expected is not None:
                if not isinstance(expected, list):
                    error(f'{label}/{dk}应为字段名数组或null')
                elif set(expected) != {f.get('name') for f in fields if f.get(fk) == '是'}:
                    error(f'{label}/{dk}与字段主键标记不一致')
        if data.get('mode') == 'requirements' and model.get('development_status') == '已完成':
            error(f'{label}仅需求场景不能标记开发已完成')
    concrete_tables = [t for t in tables if t and t != '待确认']
    if len(concrete_tables) != len(set(concrete_tables)):
        error('模型实体表名重复')
    if data.get('mode') == 'rewrite' and not data.get('migration'):
        error('rewrite必须提供逐项migration，包含原位置、目标、状态、说明')
    for item in data.get('migration', []):
        if not all(item.get(k) for k in ['source', 'target', 'status', 'note']):
            error('migration缺少source/target/status/note')
        if item.get('status') not in ('已映射', '待确认', '原样保留', '不适用'):
            error('migration.status非法')
    return errors, warnings


def table_sheet(wb, name, headers, rows, widths=None):
    ws = wb.create_sheet(name)
    for col, value in enumerate(headers, 1):
        c = put(ws, 1, col, value)
        c.fill = PatternFill('solid', fgColor='17365D')
        c.font = Font(name='微软雅黑', color='FFFFFF', bold=True)
    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            put(ws, r, c, value)
    for col in range(1, len(headers)+1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, 30)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = '1:1'
    return ws


def render(data, output, template=TEMPLATE):
    errors, warnings = validate_contract(data)
    if errors:
        raise ValueError('\n'.join(errors))
    output = Path(output)
    if output.exists():
        raise ValueError(f'拒绝覆盖：{output}')
    if output.suffix.lower() != '.xlsx':
        raise ValueError('输出必须是.xlsx')
    wb = load_workbook(template)
    if any(x not in wb.sheetnames for x in [CATALOG, DETAIL, '0 内容概览', '数据源探查']):
        raise ValueError('模板结构不匹配；其他格式须先用rewrite转成本技能规范')
    catalog, proto = wb[CATALOG], wb[DETAIL]
    design_rows, evidence_rows = [], []
    sheet_map = []
    for mi, model in enumerate(data['models'], 1):
        sheet_name = re.sub(r'[\\/*?:\[\]]', '_', model['name']).strip().strip("'")[:31] or f'模型{mi}'
        base, n = sheet_name, 1
        while sheet_name.casefold() in {x.casefold() for x in wb.sheetnames}:
            n += 1
            sheet_name = base[:27] + f'_{n}'
        ws = wb.copy_worksheet(proto)
        ws.title, ws.sheet_state = sheet_name, 'visible'
        sheet_map.append({'name': model['name'], 'sheet': sheet_name, 'table': model.get('table'), 'fields': len(model['fields'])})
        put(ws, 1, 1, f'{text(model.get("table"))}，{model["state"]}；{model.get("summary", "")}')
        put(ws, 1, 10, '版本：' + str(data.get('version', '1.0.0')))
        put(ws, 1, 12, '返回目录')
        link(ws, 'L1', CATALOG, f'E{mi+3}')
        rowno = 3
        for fi, field in enumerate(model['fields'], 1):
            start = rowno
            for mp in field['mappings']:
                for col in range(1, 13):
                    ws.cell(rowno, col)._style = copy(proto.cell(3, col)._style)
                values = [field['category'], field['subcategory'], fi, field['name'], field['type'], field['description'],
                          mp['source_field'], mp['source_table'], field['technical_key'], field['business_key'],
                          mp['expression'], f'[{field["state"]}/{mp["state"]}] 分支：{text(mp["branch"])}；条件：{text(mp["condition"])}；{field.get("note", "")}']
                for col, value in enumerate(values, 1): put(ws, rowno, col, value)
                refs = list(dict.fromkeys(field['evidence'] + mp['evidence']))
                ws.cell(rowno, 12).comment = Comment('证据：' + ', '.join(refs), 'E05')
                ws.row_dimensions[rowno].height = min(120, max(30, 15 * (1 + max(len(text(v))//45 for v in values))))
                if field['state'] != '已确认' or mp['state'] != '已确认':
                    for col in [4, 5, 12]: ws.cell(rowno, col).fill = PatternFill('solid', fgColor='FFF2CC')
                rowno += 1
            if rowno-start > 1:
                for col in [1, 2, 3, 4, 5, 6, 9, 10]:
                    ws.merge_cells(start_row=start, end_row=rowno-1, start_column=col, end_column=col)
        ws.freeze_panes = 'G3'
        ws.auto_filter.ref = f'A2:L{rowno-1}'
        ws.print_title_rows = '1:2'
        ws.print_area = f'A1:L{rowno-1}'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        for col in ['I', 'J']:
            dv = DataValidation(type='list', formula1='"是,否,待确认,不适用"')
            ws.add_data_validation(dv); dv.add(f'{col}3:{col}{rowno-1}')
        d = model['design']
        cr = mi + 3
        values = [model.get('id', f'MX{mi:05}'), model['domain'], model.get('schema'), model['layer'], model['name'],
                  model.get('designer'), model.get('design_date'), model.get('review_required', '是'),
                  model.get('development_status', '待评审'), '查看设计', model.get('developer'), model['table'],
                  model.get('start_date'), model.get('end_date'), model.get('note', model['state']),
                  model.get('business_unit'), model['domain'], model.get('business_process'), model.get('model_type'),
                  d['grain'], d['partition'], d['incremental_field'], d['frequency'], d['schedule']]
        for col, value in enumerate(values, 1):
            catalog.cell(cr, col)._style = copy(catalog.cell(4, col)._style)
            put(catalog, cr, col, value)
        link(catalog, f'E{cr}', sheet_name); link(catalog, f'J{cr}', sheet_name)
        for key in d:
            design_rows.append([model['name'], key, text(d[key]), model['state'], ', '.join(model['evidence'])])
    for e in data['evidence']:
        evidence_rows.append([e['id'], e['state'], e['claim'], e['source'], e['locator'], e.get('limitation', '')])
    table_sheet(wb, '设计说明', ['模型', '设计项', '内容', '模型状态', '证据编号'], design_rows, {3: 80})
    table_sheet(wb, '证据与待确认', ['证据编号', '状态', '结论或问题', '来源', '位置', '范围与限制'], evidence_rows, {3: 60, 4: 45, 6: 60})
    if data.get('migration'):
        table_sheet(wb, '翻写映射', ['原位置', '目标位置/语义', '处理状态', '说明'],
                    [[m[k] for k in ['source', 'target', 'status', 'note']] for m in data['migration']], {1: 50, 2: 50, 4: 70})
    for i, src in enumerate(data.get('source_checks', []), 3):
        for j, value in enumerate([i-2, src.get('name'), src.get('table'), src.get('business_key'), src.get('result', '未执行'), src.get('sql'), src.get('scope')], 2):
            put(wb['数据源探查'], i, j, value)
    for i, metric in enumerate(data.get('metrics', []), 2):
        for j, key in enumerate(['name', 'period', 'mart', 'chinese_name', 'english_name', 'algorithm'], 1):
            put(wb['1.2 OneData指标录入'], i, j, metric.get(key))
    if data.get('metrics'): wb['1.2 OneData指标录入'].sheet_state = 'visible'
    put(wb['0 内容概览'], 1, 2, data['title'])
    catalog.freeze_panes = 'F4'; catalog.auto_filter.ref = f'A3:X{len(data["models"])+3}'
    catalog.print_area = catalog.auto_filter.ref
    wb.active = wb.sheetnames.index(CATALOG)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    report = validate_workbook(output, data, sheet_map)
    report.update({'mode': data['mode'], 'warnings': warnings, 'models': sheet_map,
                   'semantic_review': '需AI逐项复核；结构通过不代表已评审或已上线',
                   'source_contract_sha256': hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()})
    return report


def validate_workbook(path, data=None, sheet_map=None):
    wb = load_workbook(path, data_only=False)
    errors = []
    for s in wb:
        for row in s:
            for c in row:
                if c.data_type == 'f' or c.data_type == 'e':
                    errors.append(f'{s.title}!{c.coordinate}存在未授权公式/Excel错误')
                if c.hyperlink and c.hyperlink.location:
                    loc = c.hyperlink.location
                    match = re.fullmatch(r"'((?:[^']|'')+)'!([A-Z]+[1-9][0-9]*)", loc)
                    if not match or match[1].replace("''", "'") not in wb.sheetnames:
                        errors.append(f'{s.title}!{c.coordinate}内部链接无效')
    if data:
        for model, info in zip(data['models'], sheet_map):
            ws = wb[info['sheet']]
            actual = [ws.cell(r, 4).value for r in range(3, ws.max_row+1) if ws.cell(r, 3).value is not None]
            if actual != [text(f['name']) for f in model['fields']]:
                errors.append(f'{ws.title}输出字段顺序/数量与合同不一致')
            if ws['A1'].value.split('，')[0] != text(model['table']):
                errors.append(f'{ws.title}表名与目录不一致')
            rowno = 3
            for f in model['fields']:
                for mp in f['mappings']:
                    for col, key in [(7, 'source_field'), (8, 'source_table'), (11, 'expression')]:
                        if ws.cell(rowno,col).value != text(mp[key]): errors.append(f'{ws.title}!{rowno}来源/表达式丢失')
                    rowno += 1
    return {'structural_pass': not errors, 'errors': errors, 'file': str(path), 'sheets': len(wb.sheetnames)}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest='action', required=True)
    a = sub.add_parser('inspect'); a.add_argument('--input', required=True); a.add_argument('--output', required=True)
    a = sub.add_parser('render'); a.add_argument('--input', required=True); a.add_argument('--output', required=True)
    a.add_argument('--template', default=str(TEMPLATE)); a.add_argument('--report', required=True)
    a = sub.add_parser('validate'); a.add_argument('--input', required=True); a.add_argument('--report', required=True)
    args = p.parse_args()
    try:
        if args.action == 'inspect': result = inspect(args.input, args.output)
        else:
            if Path(args.report).exists(): raise ValueError('report已存在，拒绝覆盖')
            if args.action == 'render':
                data = json.loads(Path(args.input).read_text(encoding='utf-8-sig'))
                result = render(data, args.output, args.template)
            else: result = validate_workbook(args.input)
            dump(args.report, result)
        print(json.dumps(result, ensure_ascii=False))
        return 0 if result.get('structural_pass', True) else 1
    except (ValueError, KeyError, TypeError, OSError) as exc:
        print(json.dumps({'error': str(exc)}, ensure_ascii=False))
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
