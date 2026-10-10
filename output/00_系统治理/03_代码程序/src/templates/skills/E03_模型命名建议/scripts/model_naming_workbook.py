#!/usr/bin/env python3
from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - runtime environment guard
    print("Missing dependency: openpyxl. Install it before using this script.", file=sys.stderr)
    raise SystemExit(2) from exc


FIELD_NAME_HEADERS = {"字段名", "字段英文名", "英文名", "字段编码", "字段code", "字段代码"}
REFERENCE_TABLE_HEADERS = {"参考库表", "参考表", "参考表名", "来源表", "来源库表", "源表", "源库表", "原表", "原始表"}
REFERENCE_FIELD_HEADERS = {"参考字段", "参考字段名", "来源字段", "来源字段名", "源字段", "源字段名", "原字段", "原字段名"}
RECOMMEND_HEADER = "推荐字段名"
RECOMMEND_BASIS_HEADER = "推荐依据"
REMARK_HEADERS = {"备注", "说明", "字段说明", "描述"}
HEADER_SCAN_ROWS = 50
MAX_SCAN_COLS = 80
BASIS_REQUIRED_LABELS = ("映射匹配", "规则匹配")
BASIS_FIELD_KEYS = ("recommendation_basis", "basis", "reason", "rationale", "source", "evidence")
TABLE_BASIS_KEYS = (
    "table_recommendation_basis",
    "recommendation_basis",
    "table_basis",
    "basis",
    "reason",
    "rationale",
    "source",
    "evidence",
)
MAPPING_MATCH_KEYS = ("mapping_match", "standard_mapping_match", "standard_match", "standard_library_match", "映射匹配")
TABLE_MAPPING_MATCH_KEYS = (
    "table_mapping_match",
    "mapping_match",
    "standard_mapping_match",
    "standard_match",
    "standard_library_match",
    "映射匹配",
)
RULE_MATCH_KEYS = ("rule_match", "naming_rule_match", "rule_evidence", "naming_rule", "规则匹配")
TABLE_RULE_MATCH_KEYS = ("table_rule_match", "rule_match", "naming_rule_match", "rule_evidence", "naming_rule", "规则匹配")
ORIGINAL_REFERENCE_KEYS = ("original_reference", "original_field_reference", "source_field_reference", "原字段引用")
AI_REASONING_KEYS = ("ai_reasoning", "semantic_reasoning", "semantic_basis", "AI推理")
CONCLUSION_KEYS = ("conclusion", "result", "结论")
CANDIDATE_NAME_KEYS = (
    "standard_candidate_name",
    "standard_field_name",
    "standard_recommended_name",
    "candidate_name",
    "mapping_candidate",
    "rule_candidate",
    "标准候选",
)
TABLE_CANDIDATE_NAME_KEYS = (
    "table_standard_candidate_name",
    "standard_candidate_name",
    "standard_table_name",
    "standard_recommended_name",
    "candidate_name",
    "mapping_candidate",
    "rule_candidate",
    "标准候选",
)
DECISION_KEYS = ("decision", "naming_decision", "recommendation_decision", "决策")
TABLE_DECISION_KEYS = ("table_decision", "decision", "naming_decision", "recommendation_decision", "决策")
MATCH_LEVEL_KEYS = ("match_level", "standard_match_level", "mapping_level", "match_type", "匹配层级")
TABLE_MATCH_LEVEL_KEYS = ("table_match_level", "match_level", "standard_match_level", "mapping_level", "match_type", "匹配层级")
KEEP_REASON_KEYS = ("keep_reason", "exception_reason", "confirm_reason", "decision_reason", "保留原因", "待确认原因")
TABLE_KEEP_REASON_KEYS = (
    "table_keep_reason",
    "keep_reason",
    "exception_reason",
    "confirm_reason",
    "decision_reason",
    "保留原因",
    "待确认原因",
)
VALID_FIELD_MATCH_LEVELS = {"exact_field", "term", "rule_only", "none"}
VALID_TABLE_MATCH_LEVELS = {"exact_table", "term", "rule_only", "none"}
VALID_DECISIONS = {"rename", "keep"}
PENDING_CONFIRMATION = "待确认"
DECISION_REASON_LABEL = "保留/修订原因"
LONG_TOKEN_MIN_LENGTH = 7
LONG_TOKEN_SUGGESTION_LABEL = "补充词根建议"
TERM_ROOT_SUGGESTION_KEYS = ("term_root_suggestions", "root_suggestions", "补充词根建议")
LONG_TOKEN_EXCEPTION_KEYS = (
    "long_token_exceptions",
    "allowed_long_tokens",
    "standard_long_tokens",
    "long_token_keep_reasons",
)
QUALITY_REPORT_FILENAME = "quality_report.md"
TOKEN_PATTERN = re.compile(r"[A-Za-z][A-Za-z0-9]*")
STANDARD_RULE_SHEET_TITLE = "标准规则库"
STANDARD_MAPPING_SHEET_TITLE = "标准映射库"
STANDARD_RULE_SUGGESTION_SHEET_TITLE = "标准规则库_建议"
STANDARD_MAPPING_SUGGESTION_SHEET_TITLE = "标准映射库_建议"
STANDARD_RULE_SUGGESTION_KEYS = ("standard_rule_suggestions", "rule_suggestions", "标准规则建议")
STANDARD_MAPPING_SUGGESTION_KEYS = ("standard_mapping_suggestions", "mapping_suggestions", "标准映射建议")
FIELD_RULE_SUGGESTION_KEYS = ("standard_rule_suggestions", "rule_suggestions", "标准规则建议")
FIELD_MAPPING_SUGGESTION_KEYS = ("standard_mapping_suggestions", "mapping_suggestions", "标准映射建议")
INVALID_SHEET_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")
MAPPING_TERM_HEADER = "中文名称"
MAPPING_ROOT_HEADER = "标准词根"
MAPPING_REQUIRED_HEADERS = ("业务域", "中文名称", "标准词根", "释义", "适用场景", "示例中文名称", "示例")
MAPPING_PLACEHOLDER_PATTERNS = (
    "本轮标准库为空",
    "标准库为空",
    "标准库未抽取到",
    "未抽取到有效数据行",
    "待标准库评审",
    "待评审",
)


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def compact(value: Any) -> str:
    return "".join(text(value).split()).lower()


def truncate(value: Any, limit: int = 240) -> str:
    raw = text(value)
    return raw if len(raw) <= limit else raw[:limit] + "..."


def english_tokens(value: Any) -> list[str]:
    return [match.group(0).lower() for match in TOKEN_PATTERN.finditer(text(value))]


def long_tokens_from_names(*names: Any) -> list[str]:
    tokens = {
        token
        for name in names
        for token in english_tokens(name)
        if len(token) > LONG_TOKEN_MIN_LENGTH
    }
    return sorted(tokens)


def ensure_xlsx(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"Only .xlsx is supported: {path}")
    if not path.exists():
        raise FileNotFoundError(path)


def is_field_name_header(value: Any) -> bool:
    normalized = compact(value)
    if normalized == compact(RECOMMEND_HEADER):
        return False
    return normalized in {compact(item) for item in FIELD_NAME_HEADERS}


def is_remark_header(value: Any) -> bool:
    return compact(value) in {compact(item) for item in REMARK_HEADERS}


def find_header_row_and_field_col(ws: Any) -> tuple[int, int] | None:
    max_row = min(ws.max_row or 0, HEADER_SCAN_ROWS)
    max_col = min(ws.max_column or 0, MAX_SCAN_COLS)
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            if is_field_name_header(ws.cell(row, col).value):
                return row, col
    return None


def header_cells(ws: Any, header_row: int) -> list[dict[str, Any]]:
    headers: list[dict[str, Any]] = []
    for col in range(1, min(ws.max_column or 0, MAX_SCAN_COLS) + 1):
        value = text(ws.cell(header_row, col).value)
        if value:
            headers.append({"col": col, "header": value})
    return headers


def find_header_col(ws: Any, header_row: int, predicate: Any) -> int | None:
    for col in range(1, ws.max_column + 1):
        if predicate(ws.cell(header_row, col).value):
            return col
    return None


def row_has_values(ws: Any, row: int, headers: list[dict[str, Any]]) -> bool:
    return any(text(ws.cell(row, item["col"]).value) for item in headers)


def field_rows(ws: Any, header_row: int, field_col: int, headers: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    blank_streak = 0
    for row in range(header_row + 1, ws.max_row + 1):
        if not row_has_values(ws, row, headers):
            blank_streak += 1
            if blank_streak >= 20:
                break
            continue
        blank_streak = 0
        current_name = text(ws.cell(row, field_col).value)
        values = {
            item["header"]: truncate(ws.cell(row, item["col"]).value)
            for item in headers
            if text(ws.cell(row, item["col"]).value)
        }
        rows.append({"row": row, "current_field_name": current_name, "values": values})
    return rows


def detect_tabular_header_row(ws: Any) -> int | None:
    max_row = min(ws.max_row or 0, 20)
    max_col = min(ws.max_column or 0, MAX_SCAN_COLS)
    best_row = None
    best_count = 0
    for row in range(1, max_row + 1):
        count = sum(1 for col in range(1, max_col + 1) if text(ws.cell(row, col).value))
        if count > best_count:
            best_row = row
            best_count = count
    return best_row if best_count >= 2 else None


def extract_standard_sheet(ws: Any, max_rows: int) -> dict[str, Any] | None:
    header_row = detect_tabular_header_row(ws)
    if not header_row:
        return None
    headers = header_cells(ws, header_row)
    records: list[dict[str, Any]] = []
    for row in range(header_row + 1, min(ws.max_row, header_row + max_rows) + 1):
        record = {
            item["header"]: truncate(ws.cell(row, item["col"]).value)
            for item in headers
            if text(ws.cell(row, item["col"]).value)
        }
        if record:
            records.append({"row": row, "values": record})
    return {"sheet": ws.title, "header_row": header_row, "headers": headers, "records": records}


def extract_context(args: argparse.Namespace) -> int:
    standard_path = Path(args.standard_lib).resolve()
    model_path = Path(args.model).resolve()
    out_path = Path(args.out).resolve()
    ensure_xlsx(standard_path)
    ensure_xlsx(model_path)

    model_wb = load_workbook(model_path, data_only=True)
    standard_wb = load_workbook(standard_path, data_only=True)

    model_sheets: list[dict[str, Any]] = []
    for ws in model_wb.worksheets:
        located = find_header_row_and_field_col(ws)
        if not located:
            continue
        header_row, field_col = located
        headers = header_cells(ws, header_row)
        model_sheets.append(
            {
                "sheet": ws.title,
                "table_name_a1": truncate(ws["A1"].value),
                "header_row": header_row,
                "field_name_col": field_col,
                "headers": headers,
                "fields": field_rows(ws, header_row, field_col, headers),
            }
        )

    standard_sheets: list[dict[str, Any]] = []
    for ws in standard_wb.worksheets:
        extracted = extract_standard_sheet(ws, args.standard_max_rows)
        if extracted:
            standard_sheets.append(extracted)

    payload = {
        "model_file": str(model_path),
        "standard_lib_file": str(standard_path),
        "naming_policy": {
            "current_field_name_usage": "current_field_name 只用于定位行和对比当前状态；不得作为标准候选、标准映射命中或保留原名依据。",
            "term_definitions": "字段名/current_field_name 是模型设计当前字段名列；参考字段/reference_field 是来源表字段列，二者不是同一列。",
            "field_name_consistency_policy": "推荐顺序第一步必须按同一 Excel 中相同 current_field_name 复用已有 recommended_field_name；相同 current_field_name 的字段推荐名必须一致。第一次出现的字段名按后续规则生成推荐，后续相同字段名直接复用该推荐。",
            "reference_table_policy": "先从字段行中的参考库表、参考表、来源表、源表、原表等列识别参考库表名；若带库名或 schema 前缀，取最后一段表名判断，大小写不敏感。参考库表名规范化后以 dwd_ 为前缀时，字段推荐以审核参考字段为主，尽量与参考字段名一致，只有与标准强规则、标准映射、命名合规性或字段语义冲突时才重新设计。参考库表名不以 dwd_ 为前缀时，不把参考字段名作为保留或风格依据，完全按标准库和字段语义推荐。",
            "recommendation_order": [
                "1. 优先按相同 current_field_name 复用本 Excel 已有 recommended_field_name，字段名相同则推荐必须一致。",
                "2. 如果没有同字段名既有推荐，再判断参考库表是否为 dwd_ 前缀且参考字段非空；满足时优先参考参考字段。",
                "3. 参考库表为空、为 ods/ods_、参考字段为空，或参考库表不是 dwd_ 前缀表时，不沿用参考字段；无同字段名既有推荐时进入标准语义推荐。",
                "4. 前三步都没有可复用推荐时，根据标准库和字段语义理解推荐，并登记到字段名一致性映射中。",
            ],
            "required_decision_fields": [
                "recommended_field_name / recommended_table_name",
                "standard_candidate_name / table_standard_candidate_name",
                "match_level / table_match_level",
                "decision / table_decision",
                "mapping_match / table_mapping_match",
                "rule_match / table_rule_match",
                "recommendation_basis / table_recommendation_basis",
            ],
            "strong_hit_rule": "exact_field、exact_table 或 term 命中且标准候选不同于当前命名时，必须推荐标准候选；不能直接保留原名。不得输出待确认，证据冲突时仍需给出负责任的具体建议并在依据中说明取舍。",
        },
        "model_sheets": model_sheets,
        "standard_sheets": standard_sheets,
    }
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"context={out_path}")
    print(f"model_sheets={len(model_sheets)}")
    print(f"standard_sheets={len(standard_sheets)}")
    return 0


def empty_recommendation_bucket() -> dict[str, Any]:
    return {"table": {}, "fields": []}


def normalize_table_item(sheet_item: dict[str, Any]) -> dict[str, Any]:
    embedded = sheet_item.get("table")
    if isinstance(embedded, dict):
        table_item = dict(embedded)
        table_item.setdefault("sheet", sheet_item.get("sheet"))
        return table_item
    return sheet_item


def normalize_recommendations(data: dict[str, Any]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for sheet_item in data.get("sheets", []):
        sheet = text(sheet_item.get("sheet"))
        if not sheet:
            continue
        result.setdefault(sheet, empty_recommendation_bucket())
        result[sheet]["table"] = normalize_table_item(sheet_item)
        result[sheet]["fields"].extend(sheet_item.get("fields", []))

    for table in data.get("tables", []):
        sheet = text(table.get("sheet"))
        if sheet:
            result.setdefault(sheet, empty_recommendation_bucket())
            result[sheet]["table"] = table

    for field in data.get("fields", []):
        sheet = text(field.get("sheet"))
        if sheet:
            result.setdefault(sheet, empty_recommendation_bucket())
            result[sheet]["fields"].append(field)
    return result


def copy_cell_style(source: Any, target: Any) -> None:
    if source.has_style:
        target._style = copy.copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy.copy(source.alignment)


def ensure_recommend_col(ws: Any, header_row: int, field_col: int) -> tuple[int, bool]:
    existing = find_header_col(ws, header_row, lambda value: compact(value) == compact(RECOMMEND_HEADER))
    if existing:
        return existing, False

    insert_at = field_col + 1
    ws.insert_cols(insert_at)
    for row in range(1, ws.max_row + 1):
        copy_cell_style(ws.cell(row, insert_at - 1), ws.cell(row, insert_at))
    ws.cell(header_row, insert_at).value = RECOMMEND_HEADER
    left_letter = get_column_letter(insert_at - 1)
    new_letter = get_column_letter(insert_at)
    ws.column_dimensions[new_letter].width = ws.column_dimensions[left_letter].width
    return insert_at, True


def ensure_recommend_basis_col(ws: Any, header_row: int, recommend_col: int) -> tuple[int, bool]:
    existing = find_header_col(ws, header_row, lambda value: compact(value) == compact(RECOMMEND_BASIS_HEADER))
    if existing:
        return existing, False

    insert_at = recommend_col + 1
    ws.insert_cols(insert_at)
    for row in range(1, ws.max_row + 1):
        copy_cell_style(ws.cell(row, insert_at - 1), ws.cell(row, insert_at))
    ws.cell(header_row, insert_at).value = RECOMMEND_BASIS_HEADER
    left_letter = get_column_letter(insert_at - 1)
    new_letter = get_column_letter(insert_at)
    left_width = ws.column_dimensions[left_letter].width or 16
    ws.column_dimensions[new_letter].width = max(left_width, 36)
    return insert_at, True


def ensure_remark_col(ws: Any, header_row: int) -> int:
    existing = find_header_col(ws, header_row, is_remark_header)
    if existing:
        return existing
    new_col = ws.max_column + 1
    ws.cell(header_row, new_col).value = "备注"
    copy_cell_style(ws.cell(header_row, new_col - 1), ws.cell(header_row, new_col))
    return new_col


def recommendation_for_row(
    row: int,
    current_name: str,
    by_row: dict[int, dict[str, Any]],
    by_name: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    if row in by_row:
        return by_row[row]
    return by_name.get(compact(current_name), {})


def append_note(existing: Any, note: str) -> str:
    current = text(existing)
    if not current:
        return note
    if note in current:
        return current
    return current + "\n" + note


def append_inline_note(existing: Any, note: str) -> str:
    current = text(existing)
    if not current:
        return note
    if note in current:
        return current
    separator = "" if current.endswith((" ", "；", ";", "，", ",")) else "；"
    return current + separator + note


def basis_has_label(value: str, label: str) -> bool:
    compacted = compact(value)
    return f"{compact(label)}:" in compacted or f"{compact(label)}：" in compacted


def basis_is_complete(value: str) -> bool:
    return all(basis_has_label(value, label) for label in BASIS_REQUIRED_LABELS)


def cjk_count(value: str) -> int:
    return sum(1 for char in value if "\u4e00" <= char <= "\u9fff")


def basis_is_chinese(value: str) -> bool:
    compacted = compact(value)
    if not compacted:
        return False
    count = cjk_count(value)
    return count >= 20 and count / max(len(compacted), 1) >= 0.18


def format_basis_value(value: Any) -> str:
    if isinstance(value, dict):
        parts = [f"{text(key)}={format_basis_value(val)}" for key, val in value.items() if text(val)]
        return "，".join(part for part in parts if part)
    if isinstance(value, list):
        parts = [format_basis_value(item) for item in value if text(item)]
        return "；".join(part for part in parts if part)
    return text(value)


def first_basis_value(item: dict[str, Any], keys: tuple[str, ...]) -> str:
    for key in keys:
        if key in item:
            value = format_basis_value(item.get(key))
            if value:
                return value
    return ""


def canonical_decision(value: Any) -> str:
    raw = compact(value).replace("-", "_")
    aliases = {
        "rename": "rename",
        "renamed": "rename",
        "change": "rename",
        "改名": "rename",
        "重命名": "rename",
        "推荐标准名": "rename",
        "采用标准名": "rename",
        "keep": "keep",
        "kept": "keep",
        "retain": "keep",
        "保留": "keep",
        "保留原名": "keep",
        "沿用": "keep",
        "沿用原名": "keep",
    }
    return aliases.get(raw, raw if raw in VALID_DECISIONS else "")


def canonical_match_level(value: Any, entity_type: str) -> str:
    raw = compact(value).replace("-", "_")
    exact_level = "exact_table" if entity_type == "table" else "exact_field"
    aliases = {
        "exact": exact_level,
        "strong": exact_level,
        "strong_hit": exact_level,
        "强匹配": exact_level,
        "精准命中": exact_level,
        "精确命中": exact_level,
        "直接命中": exact_level,
        "标准字段命中": "exact_field",
        "标准表名命中": "exact_table",
        "exactfield": "exact_field",
        "exacttable": "exact_table",
        "term": "term",
        "term_hit": "term",
        "词根命中": "term",
        "术语命中": "term",
        "语义命中": "term",
        "近似命中": "term",
        "rule": "rule_only",
        "ruleonly": "rule_only",
        "rule_only": "rule_only",
        "规则命中": "rule_only",
        "仅规则命中": "rule_only",
        "none": "none",
        "nohit": "none",
        "no_hit": "none",
        "miss": "none",
        "未命中": "none",
        "无命中": "none",
    }
    return aliases.get(raw, raw)


def same_name(left: Any, right: Any) -> bool:
    return compact(left) == compact(right)


def add_labeled_section(sections: list[str], basis: str, label: str, value: Any) -> None:
    rendered = format_basis_value(value)
    if rendered and not basis_has_label(basis, label):
        sections.append(f"{label}：{rendered}")


def compose_decision_basis(
    item: dict[str, Any],
    basis: str,
    *,
    mapping_keys: tuple[str, ...],
    rule_keys: tuple[str, ...],
    candidate_name: str,
    match_level: str,
    decision: str,
    keep_reason: str,
) -> str:
    mapping_match = first_basis_value(item, mapping_keys)
    rule_match = first_basis_value(item, rule_keys)
    ai_reasoning = first_basis_value(item, AI_REASONING_KEYS)
    conclusion = first_basis_value(item, CONCLUSION_KEYS)
    sections: list[str] = []
    if mapping_match and not basis_has_label(basis, "映射匹配"):
        sections.append(f"映射匹配：{mapping_match}")
    if rule_match and not basis_has_label(basis, "规则匹配"):
        sections.append(f"规则匹配：{rule_match}")
    if basis:
        sections.append(basis)
    add_labeled_section(sections, basis, "标准候选", candidate_name)
    add_labeled_section(sections, basis, "匹配层级", match_level)
    add_labeled_section(sections, basis, "决策", decision)
    add_labeled_section(sections, basis, DECISION_REASON_LABEL, keep_reason)
    add_labeled_section(sections, basis, LONG_TOKEN_SUGGESTION_LABEL, first_basis_value(item, TERM_ROOT_SUGGESTION_KEYS))
    if ai_reasoning and not basis_has_label(basis, "AI推理"):
        sections.append(f"AI推理：{ai_reasoning}")
    if conclusion and not basis_has_label(basis, "结论"):
        sections.append(f"结论：{conclusion}")
    return "；".join(section for section in sections if section)


def field_recommendation_payload(item: dict[str, Any]) -> dict[str, Any]:
    suggested = text(
        item.get("recommended_field_name")
        or item.get("recommendation")
        or item.get("recommended_name")
    )
    candidate_name = first_basis_value(item, CANDIDATE_NAME_KEYS)
    decision = canonical_decision(first_basis_value(item, DECISION_KEYS))
    match_level = canonical_match_level(first_basis_value(item, MATCH_LEVEL_KEYS), "field")
    keep_reason = first_basis_value(item, KEEP_REASON_KEYS)
    basis = first_basis_value(item, BASIS_FIELD_KEYS)
    basis = compose_decision_basis(
        item,
        basis,
        mapping_keys=MAPPING_MATCH_KEYS,
        rule_keys=RULE_MATCH_KEYS,
        candidate_name=candidate_name,
        match_level=match_level,
        decision=decision,
        keep_reason=keep_reason,
    )
    return {
        "name": suggested,
        "basis": basis,
        "basis_complete": basis_is_complete(basis),
        "standard_candidate_name": candidate_name,
        "decision": decision,
        "match_level": match_level,
        "keep_reason": keep_reason,
        "term_root_suggestions": first_basis_value(item, TERM_ROOT_SUGGESTION_KEYS),
        "long_token_exceptions": first_basis_value(item, LONG_TOKEN_EXCEPTION_KEYS),
    }


def table_recommendation_payload(item: Any) -> dict[str, Any]:
    if not isinstance(item, dict):
        item = {"recommended_table_name": item}
    suggested = text(
        item.get("recommended_table_name")
        or item.get("table_recommendation")
        or item.get("recommendation")
        or item.get("recommended_name")
    )
    candidate_name = first_basis_value(item, TABLE_CANDIDATE_NAME_KEYS)
    decision = canonical_decision(first_basis_value(item, TABLE_DECISION_KEYS))
    match_level = canonical_match_level(first_basis_value(item, TABLE_MATCH_LEVEL_KEYS), "table")
    keep_reason = first_basis_value(item, TABLE_KEEP_REASON_KEYS)
    basis = first_basis_value(item, TABLE_BASIS_KEYS)
    basis = compose_decision_basis(
        item,
        basis,
        mapping_keys=TABLE_MAPPING_MATCH_KEYS,
        rule_keys=TABLE_RULE_MATCH_KEYS,
        candidate_name=candidate_name,
        match_level=match_level,
        decision=decision,
        keep_reason=keep_reason,
    )
    return {
        "name": suggested,
        "basis": basis,
        "basis_complete": basis_is_complete(basis),
        "standard_candidate_name": candidate_name,
        "decision": decision,
        "match_level": match_level,
        "keep_reason": keep_reason,
        "term_root_suggestions": first_basis_value(item, TERM_ROOT_SUGGESTION_KEYS),
        "long_token_exceptions": first_basis_value(item, LONG_TOKEN_EXCEPTION_KEYS),
    }


def strong_match_levels(entity_type: str) -> set[str]:
    return {"exact_table", "term"} if entity_type == "table" else {"exact_field", "term"}


def valid_match_levels(entity_type: str) -> set[str]:
    return VALID_TABLE_MATCH_LEVELS if entity_type == "table" else VALID_FIELD_MATCH_LEVELS


def token_text_contains(value: Any, token: str) -> bool:
    normalized = text(value).lower()
    return token.lower() in english_tokens(normalized) or token.lower() in normalized


def suggestion_roots(value: Any) -> list[str]:
    raw = format_basis_value(value)
    roots = {
        match.group(1).lower()
        for match in re.finditer(r"(?:->|→)\s*([A-Za-z][A-Za-z0-9_]*)", raw)
    }
    return sorted(roots)


def name_contains_root(name: Any, root: str) -> bool:
    root_tokens = english_tokens(root)
    if not root_tokens:
        return False
    normalized_name = "_" + "_".join(english_tokens(name)) + "_"
    normalized_root = "_" + "_".join(root_tokens) + "_"
    return normalized_root in normalized_name


def payload_names_contain_root(payload: dict[str, Any], root: str) -> bool:
    return any(
        name_contains_root(payload.get(key), root)
        for key in ("name", "standard_candidate_name")
    )


def unapplied_term_root_suggestions(payload: dict[str, Any]) -> list[str]:
    roots = suggestion_roots(payload.get("term_root_suggestions"))
    return [root for root in roots if not payload_names_contain_root(payload, root)]


def basis_marks_standard_long_token(basis: str, token: str) -> bool:
    escaped = re.escape(token.lower())
    normalized = text(basis).lower()
    patterns = (
        rf"(标准词根|固定词根|固定标准名|标准固定词根|标准库).{{0,40}}\b{escaped}\b",
        rf"\b{escaped}\b.{{0,24}}(标准词根|固定词根|固定标准名|标准固定词根)",
    )
    return any(re.search(pattern, normalized) for pattern in patterns)


def long_token_reviewed(payload: dict[str, Any], token: str) -> bool:
    basis = text(payload.get("basis"))
    if token_text_contains(payload.get("long_token_exceptions"), token):
        return True
    return basis_marks_standard_long_token(basis, token)


def unreviewed_long_tokens(payload: dict[str, Any]) -> list[str]:
    tokens = long_tokens_from_names(payload.get("name"), payload.get("standard_candidate_name"))
    return [token for token in tokens if not long_token_reviewed(payload, token)]


def validate_decision_payload(
    payload: dict[str, Any],
    label: str,
    current_name: str = "",
    entity_type: str = "field",
) -> list[str]:
    errors: list[str] = []
    recommended_name = text(payload.get("name"))
    if not recommended_name:
        return errors

    if not payload.get("basis_complete"):
        errors.append(f"{label}：推荐依据必须包含 `映射匹配：...` 和 `规则匹配：...`")

    decision = text(payload.get("decision"))
    match_level = text(payload.get("match_level"))
    candidate_name = text(payload.get("standard_candidate_name"))
    keep_reason = text(payload.get("keep_reason"))
    basis = text(payload.get("basis"))
    valid_levels = valid_match_levels(entity_type)
    strong_levels = strong_match_levels(entity_type)

    if PENDING_CONFIRMATION in {recommended_name, candidate_name, decision} or PENDING_CONFIRMATION in basis:
        errors.append(f"{label}：不允许输出 `{PENDING_CONFIRMATION}`，必须给出具体且可负责的推荐名")
    if not basis_is_chinese(basis):
        errors.append(f"{label}：推荐依据必须使用中文正文，不能只写英文说明或英文规则描述")
    if not candidate_name:
        errors.append(f"{label}：必须提供具体标准候选名 `standard_candidate_name`，不能留空")
    if decision not in VALID_DECISIONS:
        errors.append(f"{label}：缺少合法决策 `decision`，只能是 rename / keep")
    if match_level not in valid_levels:
        errors.append(f"{label}：缺少合法命中层级 `match_level`，字段可用 exact_field/term/rule_only/none，表名可用 exact_table/term/rule_only/none")
    if match_level in strong_levels and not candidate_name:
        errors.append(f"{label}：标准强命中或术语命中时必须提供标准候选名 `standard_candidate_name`")
    if entity_type == "field":
        unapplied_roots = unapplied_term_root_suggestions(payload)
        if unapplied_roots:
            roots = "、".join(unapplied_roots)
            errors.append(
                f"{label}：`term_root_suggestions` / `{LONG_TOKEN_SUGGESTION_LABEL}` 中的缩写词根 `{roots}` "
                f"必须体现在推荐名 `{recommended_name}` 或标准候选 `{candidate_name}` 中；"
                "若决定保留长 token，请改用 `long_token_exceptions` / `standard_long_tokens` 说明例外"
            )
        missing_long_token_review = unreviewed_long_tokens(payload)
        if missing_long_token_review:
            tokens = "、".join(missing_long_token_review)
            errors.append(
                f"{label}：推荐名或标准候选中存在超过 {LONG_TOKEN_MIN_LENGTH} 个字母且未复核的英文 token `{tokens}`；"
                f"必须把 `{LONG_TOKEN_SUGGESTION_LABEL}：建议新增 中文词条->缩写词根` 中的缩写词根用于推荐名/标准候选，"
                "或在 `long_token_exceptions` / `standard_long_tokens` 中明确其为标准库允许的长词根"
            )

    if decision == "keep" and not keep_reason:
        errors.append(f"{label}：decision=keep 时必须提供 keep_reason / decision_reason，说明保留依据")

    if not current_name:
        return errors

    candidate_differs = bool(candidate_name) and not same_name(candidate_name, current_name)
    recommended_is_current = same_name(recommended_name, current_name)
    recommended_is_candidate = bool(candidate_name) and same_name(recommended_name, candidate_name)

    if decision == "rename":
        if recommended_is_current:
            errors.append(f"{label}：decision=rename 时推荐名不能仍等于当前名")
        if candidate_name and not recommended_is_candidate:
            errors.append(f"{label}：decision=rename 时推荐名必须等于标准候选名 `{candidate_name}`")
    elif decision == "keep":
        if not recommended_is_current:
            errors.append(f"{label}：decision=keep 时推荐名必须等于当前名 `{current_name}`")
        if match_level in strong_levels and candidate_differs:
            errors.append(
                f"{label}：已命中标准候选 `{candidate_name}`，不能直接保留当前名；必须推荐标准候选或给出另一个具体负责的候选名"
            )

    return errors


def recommendation_validation_errors(recommendations: dict[str, dict[str, Any]]) -> list[str]:
    errors: list[str] = []
    for sheet_name, sheet_recs in recommendations.items():
        table_payload = table_recommendation_payload(sheet_recs.get("table", {}))
        if sheet_recs.get("fields") and not table_payload["name"]:
            errors.append(f"{sheet_name} table：缺少表名推荐，E03 v1.5.2 要求表名与字段同样提供决策依据")
        errors.extend(validate_decision_payload(table_payload, f"{sheet_name} table", entity_type="table"))

        for item in sheet_recs.get("fields", []):
            payload = field_recommendation_payload(item)
            row = text(item.get("row")) or "?"
            current = text(item.get("current_field_name") or item.get("field_name")) or "?"
            errors.extend(validate_decision_payload(payload, f"{sheet_name} row {row} field {current}", current, "field"))
    return errors


def validate_recommendations(recommendations: dict[str, dict[str, Any]]) -> None:
    errors = recommendation_validation_errors(recommendations)
    if errors:
        sample = "；".join(errors[:10])
        more = "" if len(errors) <= 10 else f"；另有 {len(errors) - 10} 条"
        raise ValueError(
            "命名建议决策不完整或不可控：每条字段和表名推荐必须包含映射匹配、规则匹配、标准候选、命中层级和决策动作。"
            f"请补齐 recommendations.json 后重试。问题位置：{sample}{more}"
        )


def recommendation_maps(sheet_recs: dict[str, Any]) -> tuple[dict[int, dict[str, Any]], dict[str, dict[str, Any]]]:
    by_row: dict[int, dict[str, Any]] = {}
    by_name: dict[str, dict[str, Any]] = {}
    for item in sheet_recs.get("fields", []):
        payload = field_recommendation_payload(item)
        if not payload["name"]:
            continue
        if item.get("row") is not None:
            by_row[int(item["row"])] = payload
        current = text(item.get("current_field_name") or item.get("field_name"))
        if current:
            by_name[compact(current)] = payload
    return by_row, by_name


def value_for_headers(values: dict[str, Any], headers: set[str]) -> str:
    normalized_headers = {compact(header) for header in headers}
    for header, value in values.items():
        if compact(header) in normalized_headers:
            return text(value)
    return ""


def reference_table_for_row(item: dict[str, Any]) -> str:
    values = item.get("values")
    return value_for_headers(values if isinstance(values, dict) else {}, REFERENCE_TABLE_HEADERS)


def reference_field_for_row(item: dict[str, Any]) -> str:
    values = item.get("values")
    return value_for_headers(values if isinstance(values, dict) else {}, REFERENCE_FIELD_HEADERS)


def normalized_table_name(value: Any) -> str:
    raw = text(value).strip("`\"'[]")
    if not raw:
        return ""
    parts = [part for part in re.split(r"[.。/\\]", raw) if text(part)]
    return text(parts[-1]).strip("`\"'[]").lower() if parts else raw.lower()


def reference_mode_for_row(item: dict[str, Any]) -> str:
    table_name = normalized_table_name(reference_table_for_row(item))
    reference_field = reference_field_for_row(item)
    if table_name.startswith("dwd_") and reference_field:
        return "dwd_reference"
    if not table_name or table_name == "ods" or table_name.startswith("ods_") or not reference_field:
        return "field_name_consistency_or_standard"
    return "standard_semantic"


def field_name_consistency_records(wb: Any, recommendations: dict[str, dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    records: dict[str, list[dict[str, Any]]] = {}
    for ws in wb.worksheets:
        if ws.title not in recommendations:
            continue
        located = find_header_row_and_field_col(ws)
        if not located:
            continue
        header_row, field_col = located
        headers = header_cells(ws, header_row)
        rows = field_rows(ws, header_row, field_col, headers)
        by_row, by_name = recommendation_maps(recommendations[ws.title])
        for item in rows:
            current_name = text(item.get("current_field_name"))
            if not current_name:
                continue
            payload = recommendation_for_row(item["row"], current_name, by_row, by_name)
            recommended = text(payload.get("name"))
            if not recommended:
                continue
            key = compact(current_name)
            records.setdefault(key, []).append(
                {
                    "sheet": ws.title,
                    "row": item["row"],
                    "field": current_name,
                    "recommended": recommended,
                    "reference_table": reference_table_for_row(item),
                    "reference_field": reference_field_for_row(item),
                    "reference_mode": reference_mode_for_row(item),
                }
            )
    return records


def field_name_consistency_errors(wb: Any, recommendations: dict[str, dict[str, Any]]) -> list[str]:
    errors: list[str] = []
    for records in field_name_consistency_records(wb, recommendations).values():
        if len(records) <= 1:
            continue
        recommended_keys = {compact(record["recommended"]) for record in records}
        if len(recommended_keys) <= 1:
            continue
        details = "；".join(
            f"{record['sheet']} row {record['row']} 字段名 `{record['field']}` -> 推荐 `{record['recommended']}`"
            for record in records
        )
        first = records[0]
        errors.append(
            f"同字段名推荐不一致：字段名 `{first['field']}` 在同一 Excel 多个 sheet 中出现，"
            f"必须复用第一次推荐字段名 `{first['recommended']}`。当前明细：{details}"
        )
    return errors


def collect_quality_report(wb: Any, recommendations: dict[str, dict[str, Any]]) -> dict[str, Any]:
    report: dict[str, Any] = {
        "sheet_count": 0,
        "field_count": 0,
        "table_count": 0,
        "rename_count": 0,
        "keep_count": 0,
        "long_token_items": [],
        "field_name_consistency_errors": [],
        "standard_mapping_suggestion_quality": None,
        "sheets": [],
        "errors": [],
    }
    model_sheet_names: set[str] = set()

    for ws in wb.worksheets:
        located = find_header_row_and_field_col(ws)
        if not located:
            continue
        model_sheet_names.add(ws.title)
        report["sheet_count"] += 1
        header_row, field_col = located
        headers = header_cells(ws, header_row)
        rows = field_rows(ws, header_row, field_col, headers)
        sheet_report = {"sheet": ws.title, "fields": len(rows), "missing_recommendations": 0}
        report["field_count"] += len(rows)

        if ws.title not in recommendations:
            sheet_report["missing_recommendations"] = len(rows)
            report["errors"].append(f"{ws.title}：模型 sheet 缺少 recommendations.json 推荐项")
            report["sheets"].append(sheet_report)
            continue

        sheet_recs = recommendations[ws.title]
        table_payload = table_recommendation_payload(sheet_recs.get("table", {}))
        if table_payload["name"]:
            report["table_count"] += 1

        by_row, by_name = recommendation_maps(sheet_recs)
        for item in rows:
            payload = recommendation_for_row(item["row"], item["current_field_name"], by_row, by_name)
            if not payload.get("name"):
                sheet_report["missing_recommendations"] += 1
                report["errors"].append(f"{ws.title} row {item['row']} field {item['current_field_name'] or '?'}：缺少字段推荐")
                continue

            decision = text(payload.get("decision"))
            if decision == "rename":
                report["rename_count"] += 1
            elif decision == "keep":
                report["keep_count"] += 1

            tokens = long_tokens_from_names(payload.get("name"), payload.get("standard_candidate_name"))
            unapplied_roots = unapplied_term_root_suggestions(payload)
            if tokens:
                missing_review = [token for token in tokens if not long_token_reviewed(payload, token)]
                report["long_token_items"].append(
                    {
                        "sheet": ws.title,
                        "row": item["row"],
                        "field": item["current_field_name"],
                        "recommended": payload.get("name", ""),
                        "tokens": tokens,
                        "missing_review": missing_review,
                    }
                )
                if missing_review:
                    report["errors"].append(
                        f"{ws.title} row {item['row']} field {item['current_field_name'] or '?'}："
                        f"长 token 未补充词根建议或例外说明：{', '.join(missing_review)}"
                    )
            if unapplied_roots:
                report["errors"].append(
                    f"{ws.title} row {item['row']} field {item['current_field_name'] or '?'}："
                    f"补充词根建议未体现在推荐字段名或标准候选中：{', '.join(unapplied_roots)}"
                )

        report["sheets"].append(sheet_report)

    extra_sheets = sorted(set(recommendations) - model_sheet_names)
    for sheet_name in extra_sheets:
        report["errors"].append(f"{sheet_name}：recommendations.json 中存在模型工作簿未识别的 sheet")

    consistency_errors = field_name_consistency_errors(wb, recommendations)
    report["field_name_consistency_errors"] = consistency_errors
    report["errors"].extend(consistency_errors)

    report["quality_checks"] = {
        "field_count_matches_context": not any(sheet["missing_recommendations"] for sheet in report["sheets"]),
        "table_count_matches_context": report["table_count"] == report["sheet_count"],
        "field_name_consistency_passed": not consistency_errors,
        "long_token_review_completed": not any(item["missing_review"] for item in report["long_token_items"])
        and not any("补充词根建议未体现在" in error for error in report["errors"]),
        "no_pending_confirmation": True,
        "all_basis_has_mapping_and_rule": True,
    }
    return report


def write_quality_report(path: Path, report: dict[str, Any]) -> None:
    long_token_rows = []
    for item in report["long_token_items"]:
        long_token_rows.append(
            "| {sheet} | {row} | `{field}` | `{recommended}` | `{tokens}` | {status} |".format(
                sheet=item["sheet"],
                row=item["row"],
                field=item["field"],
                recommended=item["recommended"],
                tokens=", ".join(item["tokens"]),
                status="未完成" if item["missing_review"] else "已复核",
            )
        )
    if not long_token_rows:
        long_token_rows.append("| - | - | - | - | - | - |")

    sheet_rows = [
        f"| {item['sheet']} | {item['fields']} | {item['missing_recommendations']} |"
        for item in report["sheets"]
    ] or ["| - | - | - |"]
    error_rows = [f"- {error}" for error in report["errors"]] or ["- 无"]
    checks = report["quality_checks"]

    mapping_quality = report.get("standard_mapping_suggestion_quality")
    if mapping_quality:
        def mapping_items(rows: list[dict[str, Any]], include_reason: bool = False) -> str:
            if not rows:
                return "| - | - | - |" if include_reason else "| - | - |"
            rendered = []
            for item in rows:
                base = f"| {item.get('term', '')} | `{item.get('root', '')}` |"
                if include_reason:
                    base += f" {item.get('reason', '')} |"
                rendered.append(base)
            return "\n".join(rendered)

        invalid_rows = []
        for item in mapping_quality["invalid_final_rows"]:
            invalid_rows.append(f"| {item['term']} | `{item['root']}` | {'；'.join(item['errors'])} |")
        if not invalid_rows:
            invalid_rows.append("| - | - | - |")

        mapping_section = f"""
## 标准映射建议质量

| 项 | 值 |
|------|------|
| 建议初稿数 | {mapping_quality['raw_count']} |
| 已过滤已有标准数 | {len(mapping_quality['filtered_existing'])} |
| 已过滤中文名冲突数 | {len(mapping_quality['filtered_term_conflicts'])} |
| 已过滤词根冲突数 | {len(mapping_quality['filtered_root_conflicts'])} |
| 最终可复制新增数 | {mapping_quality['accepted_count']} |
| 最终建议可复制 | {mapping_quality['clean']} |

### 已过滤已有标准

| 中文名称 | 标准词根 | 已有标准 |
|------|------|------|
{mapping_items(mapping_quality['filtered_existing'], include_reason=True)}

### 已过滤中文名冲突

| 中文名称 | 建议词根 | 已有标准 |
|------|------|------|
{mapping_items(mapping_quality['filtered_term_conflicts'], include_reason=True)}

### 已过滤词根冲突

| 中文名称 | 建议词根 | 已有标准 |
|------|------|------|
{mapping_items(mapping_quality['filtered_root_conflicts'], include_reason=True)}

### 最终可复制新增项

| 中文名称 | 标准词根 |
|------|------|
{mapping_items(mapping_quality['final_rows'])}

### 无法写入的最终项

| 中文名称 | 标准词根 | 问题 |
|------|------|------|
{chr(10).join(invalid_rows)}
"""
    else:
        mapping_section = """
## 标准映射建议质量

未启用：本次写回未传入标准库，无法校验 `标准映射库_建议` 是否可复制。
"""

    content = f"""# E03 命名建议质量报告

## 汇总

| 项 | 值 |
|------|------|
| 模型 sheet 数 | {report['sheet_count']} |
| 字段数 | {report['field_count']} |
| 表名推荐数 | {report['table_count']} |
| rename 数 | {report['rename_count']} |
| keep 数 | {report['keep_count']} |

## 质量门禁

| 检查项 | 结果 |
|------|------|
| 字段推荐覆盖全部模型字段 | {checks['field_count_matches_context']} |
| 表名推荐覆盖全部模型 sheet | {checks['table_count_matches_context']} |
| 同字段名推荐一致 | {checks['field_name_consistency_passed']} |
| 长 token 已补充词根建议或例外说明 | {checks['long_token_review_completed']} |
| 无待确认占位 | {checks['no_pending_confirmation']} |
| 依据包含映射匹配和规则匹配 | {checks['all_basis_has_mapping_and_rule']} |
| 标准映射建议可复制 | {checks.get('standard_mapping_suggestions_clean', True)} |

## Sheet 覆盖

| sheet | 字段数 | 缺少推荐数 |
|------|------|------|
{chr(10).join(sheet_rows)}

## 长 Token 复核

| sheet | row | 当前字段 | 推荐字段 | 长 token | 状态 |
|------|------|------|------|------|------|
{chr(10).join(long_token_rows)}

{mapping_section}

## 错误

{chr(10).join(error_rows)}
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def safe_sheet_title(title: str) -> str:
    cleaned = INVALID_SHEET_TITLE_CHARS.sub("_", text(title)) or "Sheet"
    return cleaned[:31]


def unique_sheet_title(wb: Any, title: str) -> str:
    base = safe_sheet_title(title)
    if base not in wb.sheetnames:
        return base
    for index in range(2, 1000):
        suffix = f"_{index}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        if candidate not in wb.sheetnames:
            return candidate
    raise ValueError(f"Cannot create unique sheet title for {title}")


def copy_cell(source_cell: Any, target_cell: Any, include_value: bool = True) -> None:
    if include_value:
        target_cell.value = source_cell.value
    if source_cell.has_style:
        # Copy style components instead of the private style array.  The source
        # and target sheets can belong to different workbooks, and reusing the
        # private style id can point at an alignment/fill table entry that does
        # not exist in the target workbook when openpyxl saves it.
        target_cell.font = copy.copy(source_cell.font)
        target_cell.fill = copy.copy(source_cell.fill)
        target_cell.border = copy.copy(source_cell.border)
        target_cell.alignment = copy.copy(source_cell.alignment)
        target_cell.protection = copy.copy(source_cell.protection)
    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format
    if source_cell.hyperlink:
        target_cell._hyperlink = copy.copy(source_cell.hyperlink)
    if source_cell.comment:
        target_cell.comment = copy.copy(source_cell.comment)


def copy_sheet_layout(source_ws: Any, target_ws: Any, max_row: int | None = None) -> None:
    for key, source_dim in source_ws.column_dimensions.items():
        target_dim = target_ws.column_dimensions[key]
        target_dim.width = source_dim.width
        target_dim.hidden = source_dim.hidden
        target_dim.outlineLevel = source_dim.outlineLevel
        target_dim.collapsed = source_dim.collapsed
    for key, source_dim in source_ws.row_dimensions.items():
        if max_row is not None and key > max_row:
            continue
        target_dim = target_ws.row_dimensions[key]
        target_dim.height = source_dim.height
        target_dim.hidden = source_dim.hidden
        target_dim.outlineLevel = source_dim.outlineLevel
        target_dim.collapsed = source_dim.collapsed
    target_ws.freeze_panes = source_ws.freeze_panes
    target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines


def copy_worksheet_cells(source_ws: Any, target_ws: Any, max_row: int | None = None) -> None:
    row_limit = max_row or source_ws.max_row
    max_col = source_ws.max_column
    for row in source_ws.iter_rows(min_row=1, max_row=row_limit, max_col=max_col):
        for source_cell in row:
            copy_cell(source_cell, target_ws.cell(source_cell.row, source_cell.column))
    for merged_range in source_ws.merged_cells.ranges:
        if max_row is None or merged_range.max_row <= max_row:
            target_ws.merge_cells(str(merged_range))
    copy_sheet_layout(source_ws, target_ws, max_row=max_row)


def copy_worksheet_to_workbook(source_ws: Any, target_wb: Any, title: str) -> Any:
    target_ws = target_wb.create_sheet(unique_sheet_title(target_wb, title))
    copy_worksheet_cells(source_ws, target_ws)
    target_ws.sheet_state = source_ws.sheet_state
    if source_ws.auto_filter.ref:
        target_ws.auto_filter.ref = source_ws.auto_filter.ref
    return target_ws


def first_item_value(item: dict[str, Any], keys: tuple[str, ...]) -> str:
    for key in keys:
        if key in item and text(item.get(key)):
            return text(item.get(key))
    return ""


def iter_suggestion_entries(value: Any) -> list[Any]:
    if not value:
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, dict):
        for key in ("items", "suggestions", "rows", "values"):
            nested = value.get(key)
            if isinstance(nested, list):
                return nested
        return [value]
    return [value]


def parse_term_root_pairs(value: Any) -> list[tuple[str, str]]:
    raw = format_basis_value(value)
    pairs: list[tuple[str, str]] = []
    for chunk in re.split(r"[；;\n\r]+", raw):
        match = re.search(r"(.+?)(?:->|→)\s*([A-Za-z][A-Za-z0-9_]*)", chunk)
        if not match:
            continue
        term = re.sub(r"^(建议新增|新增|建议)\s*", "", match.group(1)).strip(" ：:，,、")
        root = match.group(2).strip()
        if term and root:
            pairs.append((term, root))
    return pairs


def mapping_suggestion_from_pair(term: str, root: str, source: str = "", example: str = "") -> dict[str, Any]:
    row = {
        "业务域": "待评审",
        "中文名称": term,
        "标准词根": root,
        "释义": "由 E03 命名建议沉淀，需评审后纳入标准映射库。",
        "适用场景": f"待标准库评审；来源：{source}" if source else "待标准库评审",
        "示例中文名称": term,
        "示例": example or root,
    }
    return row


def normalize_mapping_suggestion(entry: Any, source: str = "", example: str = "") -> list[dict[str, Any]]:
    if isinstance(entry, str):
        return [mapping_suggestion_from_pair(term, root, source, example) for term, root in parse_term_root_pairs(entry)]
    if not isinstance(entry, dict):
        return []
    rows: list[dict[str, Any]] = []
    for key in TERM_ROOT_SUGGESTION_KEYS:
        rows.extend(mapping_suggestion_from_pair(term, root, source, example) for term, root in parse_term_root_pairs(entry.get(key)))
    term = first_item_value(entry, ("中文名称", "term", "chinese_name", "business_term", "name", "词条"))
    root = first_item_value(entry, ("标准词根", "standard_root", "root", "term_root", "abbr", "缩写词根"))
    if term and root:
        row = dict(entry)
        row.setdefault("中文名称", term)
        row.setdefault("标准词根", root)
        row.setdefault("业务域", first_item_value(entry, ("业务域", "domain", "business_domain")) or "待评审")
        row.setdefault("释义", first_item_value(entry, ("释义", "definition", "description", "reason")) or "由 E03 命名建议沉淀，需评审后纳入标准映射库。")
        row.setdefault("适用场景", first_item_value(entry, ("适用场景", "scenario", "usage", "source")) or (f"待标准库评审；来源：{source}" if source else "待标准库评审"))
        row.setdefault("示例中文名称", first_item_value(entry, ("示例中文名称", "example_chinese_name")) or term)
        row.setdefault("示例", first_item_value(entry, ("示例", "example", "recommended_field_name")) or example or root)
        rows.append(row)
    return rows


def normalize_rule_suggestion(entry: Any, source: str = "") -> list[dict[str, Any]]:
    if isinstance(entry, str):
        description = text(entry)
        if not description:
            return []
        return [{"作用对象": "字段名", "规则描述": description, "示例": source}]
    if not isinstance(entry, dict):
        return []
    description = first_item_value(entry, ("规则描述", "rule_description", "description", "rule", "suggestion", "建议"))
    if not description:
        return []
    row = dict(entry)
    row.setdefault("作用对象", first_item_value(entry, ("作用对象", "object", "target", "target_object")) or "字段名")
    row.setdefault("规则描述", description)
    row.setdefault("示例", first_item_value(entry, ("示例", "example")) or source)
    return [row]


def collect_raw_standard_mapping_suggestions(data: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()

    def add(row: dict[str, Any]) -> None:
        term = text(row.get(MAPPING_TERM_HEADER))
        root = text(row.get(MAPPING_ROOT_HEADER))
        if not term or not root:
            return
        key = (compact(term), compact(root))
        if key in seen:
            return
        seen.add(key)
        row["序号"] = len(rows) + 1
        rows.append(row)

    for key in STANDARD_MAPPING_SUGGESTION_KEYS:
        for entry in iter_suggestion_entries(data.get(key)):
            for row in normalize_mapping_suggestion(entry):
                add(row)

    for sheet_item in data.get("sheets", []):
        sheet_name = text(sheet_item.get("sheet"))
        for container in [sheet_item.get("table", {})] + list(sheet_item.get("fields", [])):
            if not isinstance(container, dict):
                continue
            source = f"{sheet_name} row {container.get('row')} {container.get('current_field_name') or container.get('recommended_field_name')}".strip()
            example = text(container.get("recommended_field_name") or container.get("standard_candidate_name"))
            for key in TERM_ROOT_SUGGESTION_KEYS + FIELD_MAPPING_SUGGESTION_KEYS:
                for row in normalize_mapping_suggestion(container.get(key), source=source, example=example):
                    add(row)
    return rows


def standard_mapping_records_from_workbook(standard_wb: Any) -> list[dict[str, Any]]:
    template = find_standard_template_sheet(standard_wb, STANDARD_MAPPING_SHEET_TITLE, {MAPPING_TERM_HEADER, MAPPING_ROOT_HEADER})
    if not template:
        return []
    header_row = detect_tabular_header_row(template) or 1
    headers = header_cells(template, header_row)
    records: list[dict[str, Any]] = []
    for row_num in range(header_row + 1, template.max_row + 1):
        record = {
            item["header"]: text(template.cell(row_num, item["col"]).value)
            for item in headers
            if text(template.cell(row_num, item["col"]).value)
        }
        if text(record.get(MAPPING_TERM_HEADER)) and text(record.get(MAPPING_ROOT_HEADER)):
            record["_row"] = row_num
            records.append(record)
    return records


def mapping_index(records: list[dict[str, Any]]) -> dict[str, Any]:
    pairs: dict[tuple[str, str], dict[str, Any]] = {}
    terms: dict[str, list[dict[str, Any]]] = {}
    roots: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        term = text(record.get(MAPPING_TERM_HEADER))
        root = text(record.get(MAPPING_ROOT_HEADER))
        if not term or not root:
            continue
        pair_key = (compact(term), compact(root))
        pairs[pair_key] = record
        terms.setdefault(compact(term), []).append(record)
        roots.setdefault(compact(root), []).append(record)
    return {"pairs": pairs, "terms": terms, "roots": roots}


def format_existing_mapping(records: list[dict[str, Any]]) -> str:
    parts = []
    for record in records[:5]:
        row = record.get("_row")
        term = text(record.get(MAPPING_TERM_HEADER))
        root = text(record.get(MAPPING_ROOT_HEADER))
        prefix = f"第{row}行" if row else "已有行"
        parts.append(f"{prefix}{term}->{root}")
    return "；".join(parts)


def mapping_suggestion_is_complete(row: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    for header in MAPPING_REQUIRED_HEADERS:
        if not text(row.get(header)):
            errors.append(f"缺少{header}")
    rendered = "；".join(text(row.get(header)) for header in MAPPING_REQUIRED_HEADERS)
    for pattern in MAPPING_PLACEHOLDER_PATTERNS:
        if pattern in rendered:
            errors.append(f"存在占位说明：{pattern}")
    return errors


def standard_mapping_suggestion_quality(
    raw_rows: list[dict[str, Any]],
    existing_records: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    index = mapping_index(existing_records)
    accepted: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    quality: dict[str, Any] = {
        "raw_count": len(raw_rows),
        "accepted_count": 0,
        "filtered_existing": [],
        "filtered_root_conflicts": [],
        "filtered_term_conflicts": [],
        "invalid_final_rows": [],
        "final_rows": [],
        "clean": True,
    }

    for row in raw_rows:
        term = text(row.get(MAPPING_TERM_HEADER))
        root = text(row.get(MAPPING_ROOT_HEADER))
        if not term or not root:
            continue
        pair_key = (compact(term), compact(root))
        if pair_key in seen:
            continue
        seen.add(pair_key)

        if pair_key in index["pairs"]:
            quality["filtered_existing"].append(
                {
                    "term": term,
                    "root": root,
                    "reason": format_existing_mapping([index["pairs"][pair_key]]),
                }
            )
            continue

        existing_term_rows = index["terms"].get(compact(term), [])
        if existing_term_rows:
            quality["filtered_term_conflicts"].append(
                {
                    "term": term,
                    "root": root,
                    "reason": format_existing_mapping(existing_term_rows),
                }
            )
            continue

        existing_root_rows = index["roots"].get(compact(root), [])
        if existing_root_rows:
            quality["filtered_root_conflicts"].append(
                {
                    "term": term,
                    "root": root,
                    "reason": format_existing_mapping(existing_root_rows),
                }
            )
            continue

        errors = mapping_suggestion_is_complete(row)
        if errors:
            quality["invalid_final_rows"].append({"term": term, "root": root, "errors": errors})
            continue

        clean_row = dict(row)
        clean_row["序号"] = len(accepted) + 1
        accepted.append(clean_row)
        quality["final_rows"].append({"term": term, "root": root})

    quality["accepted_count"] = len(accepted)
    quality["clean"] = not quality["invalid_final_rows"]
    return accepted, quality


def collect_standard_mapping_suggestions(
    data: dict[str, Any],
    existing_records: list[dict[str, Any]] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    raw_rows = collect_raw_standard_mapping_suggestions(data)
    return standard_mapping_suggestion_quality(raw_rows, existing_records or [])


def collect_standard_rule_suggestions(data: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()

    def add(row: dict[str, Any]) -> None:
        description = text(row.get("规则描述"))
        if not description:
            return
        key = compact(description)
        if key in seen:
            return
        seen.add(key)
        row["序号"] = len(rows) + 1
        rows.append(row)

    for key in STANDARD_RULE_SUGGESTION_KEYS:
        for entry in iter_suggestion_entries(data.get(key)):
            for row in normalize_rule_suggestion(entry):
                add(row)

    for sheet_item in data.get("sheets", []):
        sheet_name = text(sheet_item.get("sheet"))
        for container in [sheet_item.get("table", {})] + list(sheet_item.get("fields", [])):
            if not isinstance(container, dict):
                continue
            source = f"{sheet_name} row {container.get('row')} {container.get('current_field_name') or container.get('recommended_field_name')}".strip()
            for key in FIELD_RULE_SUGGESTION_KEYS:
                for entry in iter_suggestion_entries(container.get(key)):
                    for row in normalize_rule_suggestion(entry, source=source):
                        add(row)
    return rows


def find_standard_template_sheet(standard_wb: Any, preferred_title: str, required_headers: set[str]) -> Any | None:
    if preferred_title in standard_wb.sheetnames:
        return standard_wb[preferred_title]
    for ws in standard_wb.worksheets:
        header_row = detect_tabular_header_row(ws)
        if not header_row:
            continue
        headers = {item["header"] for item in header_cells(ws, header_row)}
        if required_headers.issubset(headers):
            return ws
    return None


def create_suggestion_sheet(target_wb: Any, template_ws: Any, title: str, rows: list[dict[str, Any]]) -> Any:
    header_row = detect_tabular_header_row(template_ws) or 1
    headers = header_cells(template_ws, header_row)
    target_ws = target_wb.create_sheet(unique_sheet_title(target_wb, title))
    copy_worksheet_cells(template_ws, target_ws, max_row=header_row)
    data_style_row = header_row + 1 if template_ws.max_row > header_row else header_row
    max_col = max(template_ws.max_column, max((item["col"] for item in headers), default=1))
    for offset, row_data in enumerate(rows, start=1):
        target_row = header_row + offset
        for col in range(1, max_col + 1):
            copy_cell(template_ws.cell(data_style_row, col), target_ws.cell(target_row, col), include_value=False)
        for header in headers:
            value = row_data.get(header["header"], "")
            target_ws.cell(target_row, header["col"]).value = value
    return target_ws


def append_standard_library_sheets(target_wb: Any, standard_path: Path, data: dict[str, Any]) -> dict[str, int]:
    standard_wb = load_workbook(standard_path)
    result = {"standard_sheets_copied": 0, "suggestion_sheets_written": 0, "mapping_suggestions": 0, "rule_suggestions": 0}
    try:
        for ws in standard_wb.worksheets:
            copy_worksheet_to_workbook(ws, target_wb, ws.title)
            result["standard_sheets_copied"] += 1

        rule_rows = collect_standard_rule_suggestions(data)
        existing_mapping_records = standard_mapping_records_from_workbook(standard_wb)
        mapping_rows, _ = collect_standard_mapping_suggestions(data, existing_mapping_records)
        rule_template = find_standard_template_sheet(standard_wb, STANDARD_RULE_SHEET_TITLE, {"作用对象", "规则描述"})
        mapping_template = find_standard_template_sheet(standard_wb, STANDARD_MAPPING_SHEET_TITLE, {"中文名称", "标准词根"})
        if rule_template:
            create_suggestion_sheet(target_wb, rule_template, STANDARD_RULE_SUGGESTION_SHEET_TITLE, rule_rows)
            result["suggestion_sheets_written"] += 1
            result["rule_suggestions"] = len(rule_rows)
        if mapping_template:
            create_suggestion_sheet(target_wb, mapping_template, STANDARD_MAPPING_SUGGESTION_SHEET_TITLE, mapping_rows)
            result["suggestion_sheets_written"] += 1
            result["mapping_suggestions"] = len(mapping_rows)
    finally:
        standard_wb.close()
    return result


def table_name_from_a1(value: Any) -> str:
    current = text(value)
    for marker in ("\n推荐表名：", "；推荐表名：", ";推荐表名：", " 推荐表名："):
        if marker in current:
            return current.split(marker, 1)[0].strip()
    return current


def table_recommendation_note(payload: dict[str, Any]) -> str:
    return f"推荐表名：{payload['name']}；表名推荐依据：{payload['basis']}"


def write_recommendations(args: argparse.Namespace) -> int:
    model_path = Path(args.model).resolve()
    standard_path = Path(args.standard_lib).resolve() if args.standard_lib else None
    recommendations_path = Path(args.recommendations).resolve()
    output_path = Path(args.output).resolve() if args.output else model_path.with_name(model_path.stem + "_命名建议.xlsx")
    ensure_xlsx(model_path)
    if standard_path:
        ensure_xlsx(standard_path)
    if not recommendations_path.exists():
        raise FileNotFoundError(recommendations_path)
    if output_path.exists() and not args.overwrite:
        raise FileExistsError(f"Output already exists: {output_path}. Use --overwrite or choose another --output.")

    data = json.loads(recommendations_path.read_text(encoding="utf-8"))
    recommendations = normalize_recommendations(data)
    quality_report_path = (
        Path(args.quality_report).resolve()
        if args.quality_report
        else recommendations_path.with_name(QUALITY_REPORT_FILENAME)
    )
    source_wb = load_workbook(model_path, data_only=True)
    quality_report = collect_quality_report(source_wb, recommendations)
    source_wb.close()
    quality_report["quality_checks"]["standard_mapping_suggestions_clean"] = True
    if standard_path:
        standard_wb_for_quality = load_workbook(standard_path, data_only=True)
        try:
            existing_mapping_records = standard_mapping_records_from_workbook(standard_wb_for_quality)
            _, mapping_quality = collect_standard_mapping_suggestions(data, existing_mapping_records)
        finally:
            standard_wb_for_quality.close()
        quality_report["standard_mapping_suggestion_quality"] = mapping_quality
        quality_report["quality_checks"]["standard_mapping_suggestions_clean"] = mapping_quality["clean"]
        if not mapping_quality["clean"]:
            quality_report["errors"].append("标准映射库_建议存在不可直接复制的最终项，请补齐业务域、释义、适用场景、示例并移除占位说明")
    decision_errors = recommendation_validation_errors(recommendations)
    if decision_errors:
        for error in decision_errors:
            if error not in quality_report["errors"]:
                quality_report["errors"].append(error)
        quality_report["quality_checks"]["no_pending_confirmation"] = not any(PENDING_CONFIRMATION in error for error in decision_errors)
        quality_report["quality_checks"]["long_token_review_completed"] = not any(
            "长 token" in error or "词根建议" in error
            for error in decision_errors
        )
        quality_report["quality_checks"]["all_basis_has_mapping_and_rule"] = not any(
            "推荐依据必须包含" in error or "推荐依据必须使用中文正文" in error
            for error in decision_errors
        )
    write_quality_report(quality_report_path, quality_report)
    if quality_report["errors"]:
        sample = "；".join(quality_report["errors"][:10])
        more = "" if len(quality_report["errors"]) <= 10 else f"；另有 {len(quality_report['errors']) - 10} 条"
        raise ValueError(f"E03 质量门禁失败，已生成 quality_report.md。问题位置：{sample}{more}")
    validate_recommendations(recommendations)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(model_path, output_path)

    wb = load_workbook(output_path)
    touched_sheets = 0
    touched_fields = 0
    touched_basis = 0
    standard_report = {"standard_sheets_copied": 0, "suggestion_sheets_written": 0, "mapping_suggestions": 0, "rule_suggestions": 0}

    for ws in wb.worksheets:
        if ws.title not in recommendations:
            continue
        located = find_header_row_and_field_col(ws)
        if not located:
            raise ValueError(f"Cannot find field name header in sheet: {ws.title}")
        header_row, field_col = located
        headers = header_cells(ws, header_row)
        rows = field_rows(ws, header_row, field_col, headers)
        recommend_col, _ = ensure_recommend_col(ws, header_row, field_col)
        basis_col, _ = ensure_recommend_basis_col(ws, header_row, recommend_col)

        sheet_recs = recommendations[ws.title]
        by_row, by_name = recommendation_maps(sheet_recs)

        for item in rows:
            payload = recommendation_for_row(item["row"], item["current_field_name"], by_row, by_name)
            suggested = payload.get("name", "")
            if suggested:
                decision_errors = validate_decision_payload(
                    payload,
                    f"{ws.title} row {item['row']} field {item['current_field_name'] or '?'}",
                    item["current_field_name"],
                    "field",
                )
                if decision_errors:
                    raise ValueError("；".join(decision_errors))
                ws.cell(item["row"], recommend_col).value = suggested
                basis = payload.get("basis", "")
                if not basis_is_complete(basis):
                    raise ValueError(f"Recommendation basis became incomplete while writing sheet {ws.title}, row {item['row']}.")
                ws.cell(item["row"], basis_col).value = basis
                touched_fields += 1
                touched_basis += 1

        table_payload = table_recommendation_payload(sheet_recs.get("table", {}))
        table_name = text(table_payload.get("name"))
        if table_name:
            current_table_name = table_name_from_a1(ws["A1"].value)
            decision_errors = validate_decision_payload(table_payload, f"{ws.title} table", current_table_name, "table")
            if decision_errors:
                raise ValueError("；".join(decision_errors))
            note = table_recommendation_note(table_payload)
            ws["A1"].value = append_inline_note(ws["A1"].value, note)

        touched_sheets += 1

    if standard_path:
        standard_report = append_standard_library_sheets(wb, standard_path, data)

    wb.save(output_path)
    print(f"output={output_path}")
    print(f"touched_sheets={touched_sheets}")
    print(f"touched_fields={touched_fields}")
    print(f"touched_basis={touched_basis}")
    print(f"standard_sheets_copied={standard_report['standard_sheets_copied']}")
    print(f"suggestion_sheets_written={standard_report['suggestion_sheets_written']}")
    print(f"mapping_suggestions={standard_report['mapping_suggestions']}")
    print(f"rule_suggestions={standard_report['rule_suggestions']}")
    print(f"quality_report={quality_report_path}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Extract model naming context and write naming recommendations to a copied workbook.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    extract = subparsers.add_parser("extract", help="Extract compact JSON context from standard library and model workbook.")
    extract.add_argument("--standard-lib", required=True, help="Path to standard library .xlsx")
    extract.add_argument("--model", required=True, help="Path to model design .xlsx")
    extract.add_argument("--out", required=True, help="Output context JSON path")
    extract.add_argument("--standard-max-rows", type=int, default=2000, help="Maximum rows per standard sheet")
    extract.set_defaults(func=extract_context)

    write = subparsers.add_parser("write", help="Copy model workbook and write recommendation JSON into it.")
    write.add_argument("--model", required=True, help="Path to model design .xlsx")
    write.add_argument("--standard-lib", help="Optional standard library .xlsx to copy into the output workbook and use as templates for suggestion sheets")
    write.add_argument("--recommendations", required=True, help="AI recommendations JSON path")
    write.add_argument("--output", help="Output .xlsx path; defaults to *_命名建议.xlsx next to model")
    write.add_argument("--quality-report", help="Output Markdown quality report path; defaults to quality_report.md next to recommendations")
    write.add_argument("--overwrite", action="store_true", help="Allow overwriting output path")
    write.set_defaults(func=write_recommendations)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        return args.func(args)
    except Exception as exc:
        print(f"error={exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
