#!/usr/bin/env python3
from __future__ import annotations

import argparse
from copy import copy
from datetime import date, datetime
import json
from pathlib import Path
import re
import shutil
import sys

from openpyxl import load_workbook


SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "数仓模型交付文档模板.xlsx"
DIMENSIONS = ["完整性", "唯一性", "有效性", "准确性", "一致性", "及时性"]
QUALITY_SQL_RE = re.compile(
    r"(?is)select\s+count\s*\(\s*\*\s*\)\s+as\s+quality_value\s+from\s+dc\s*;?\s*$"
)


def safe_sheet_name(prefix: str, model_name: str, suffix: str) -> str:
    forbidden = r"[]:*?/\\"
    clean = "".join("_" if char in forbidden else char for char in model_name).strip() or "模型名称"
    available = 31 - len(prefix) - len(suffix)
    return f"{prefix}{clean[:max(1, available)]}{suffix}"


def find_sheet(workbook, prefix: str):
    matches = [name for name in workbook.sheetnames if name.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"必须且只能存在一个以 {prefix} 开头的 sheet，实际为：{matches}")
    return workbook[matches[0]]


def copy_row_style(ws, source_row: int, target_row: int, start_col: int, end_col: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(start_col, end_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)


def ensure_rows(ws, start_row: int, target_count: int, style_row: int, start_col: int, end_col: int) -> None:
    required_last = start_row + target_count - 1
    while ws.max_row < required_last:
        new_row = ws.max_row + 1
        copy_row_style(ws, style_row, new_row, start_col, end_col)


def clear_range(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            ws.cell(row, col).value = None


def build_workbook(package: dict, template: Path, output: Path, overwrite: bool) -> None:
    if not template.exists():
        raise FileNotFoundError(f"模板不存在：{template}")
    if output.exists() and not overwrite:
        raise FileExistsError(f"输出已存在，需显式使用 --overwrite：{output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output)

    metadata = package.get("metadata", {})
    model = package["model"]
    fields = package.get("fields", [])
    quality = package.get("quality", [])
    tests = package.get("tests", [])
    model_name = str(model["name"]).strip()

    wb = load_workbook(output)
    ws_list = wb["01_交付表清单"]
    ws_fields = find_sheet(wb, "02_1")
    ws_quality = find_sheet(wb, "02_2")
    ws_tests = find_sheet(wb, "02_3")
    ws_fields.title = safe_sheet_name("02_1", model_name, "-字段")
    ws_quality.title = safe_sheet_name("02_2", model_name, "-质量指标")
    ws_tests.title = safe_sheet_name("02_3", model_name, "-数据自测")

    delivery_date = metadata.get("delivery_date") or date.today().isoformat()
    if isinstance(delivery_date, str):
        try:
            delivery_date = datetime.strptime(delivery_date, "%Y-%m-%d")
        except ValueError:
            pass
    owner = metadata.get("owner", model.get("owner", ""))
    ws_list["C2"] = model_name
    ws_list["G2"] = metadata.get("document_version", "V1.0")
    ws_list["I2"] = delivery_date
    ws_list["K2"] = owner
    list_values = [
        1, model_name, model.get("layer", ""), model.get("table_name", ""),
        model.get("schedule", ""), model.get("scenario", ""),
        "\n".join(model.get("sources", [])), "-", owner, model.get("remark", ""),
    ]
    for col, value in enumerate(list_values, start=2):
        ws_list.cell(5, col).value = value

    ensure_rows(ws_fields, 6, max(1, len(fields)), 6, 2, 8)
    clear_range(ws_fields, 6, ws_fields.max_row, 2, 8)
    for index, field in enumerate(fields, start=6):
        values = [
            field.get("seq", index - 5), field["cn_name"], field["name"], field["type"],
            field["description"], field["logic"], field["source"],
        ]
        for col, value in enumerate(values, start=2):
            ws_fields.cell(index, col).value = value
    ws_fields.auto_filter.ref = f"B2:H{max(5, 5 + len(fields))}"

    clear_range(ws_quality, 3, 8, 2, 13)
    for index, item in enumerate(quality, start=3):
        values = [
            item.get("seq", index - 2), f"{item['dimension']}-{model_name}", item["dimension"],
            item["object"], item["rule"], item["threshold"], item.get("period", "每日"), None,
            item.get("note", ""), item["sql"], item.get("actual", "待实测"),
            item.get("conclusion", "待实测"),
        ]
        for col, value in enumerate(values, start=2):
            ws_quality.cell(index, col).value = value
    ws_quality.auto_filter.ref = "B2:M8"

    ensure_rows(ws_tests, 3, max(1, len(tests)), 3, 2, 11)
    clear_range(ws_tests, 3, ws_tests.max_row, 2, 11)
    for index, item in enumerate(tests, start=3):
        values = [
            item.get("seq", index - 2), item["scenario"], item["method"], item["object"],
            item["expected"], item.get("actual", "待实测"), item.get("scope", ""), None,
            item.get("note", ""), item.get("sql", ""),
        ]
        for col, value in enumerate(values, start=2):
            ws_tests.cell(index, col).value = value
    ws_tests.auto_filter.ref = f"B2:K{max(3, 2 + len(tests))}"
    wb.save(output)


def validate_workbook(path: Path, allow_pending: bool) -> list[str]:
    errors: list[str] = []
    wb = load_workbook(path, data_only=False)
    try:
        fields = find_sheet(wb, "02_1")
        quality = find_sheet(wb, "02_2")
        tests = find_sheet(wb, "02_3")
    except ValueError as exc:
        wb.close()
        return [str(exc)]

    dimensions = [quality.cell(row, 4).value for row in range(3, 9)]
    if dimensions != DIMENSIONS:
        errors.append(f"六维顺序错误：{dimensions}")
    model_names = []
    for row, dimension in enumerate(DIMENSIONS, start=3):
        name = str(quality.cell(row, 3).value or "")
        if not name.startswith(f"{dimension}-"):
            errors.append(f"第{row}行指标名称不是“{dimension}-模型名称”：{name}")
        model_names.append(name.split("-", 1)[-1] if "-" in name else "")
        threshold = str(quality.cell(row, 7).value or "")
        if "quality_value" in threshold:
            errors.append(f"第{row}行业务阈值混入 quality_value 技术协议")
        sql = str(quality.cell(row, 11).value or "").strip()
        if not QUALITY_SQL_RE.search(sql):
            errors.append(f"第{row}行 SQL 未以 SELECT COUNT(*) AS quality_value FROM dc 结束")
        if not allow_pending:
            for col in (12, 13):
                value = str(quality.cell(row, col).value or "")
                if not value or "待实测" in value:
                    errors.append(f"第{row}行实际结果或结论仍待实测")
    if len(set(model_names)) > 1:
        errors.append(f"指标名称中的模型名称不一致：{model_names}")

    field_names = [
        str(fields.cell(row, 4).value)
        for row in range(6, fields.max_row + 1)
        if fields.cell(row, 4).value
    ]
    if len(field_names) != len(set(field_names)):
        errors.append("字段技术名存在重复")
    if not field_names:
        errors.append("字段页没有业务字段")

    if not allow_pending:
        for row in range(3, tests.max_row + 1):
            if tests.cell(row, 3).value:
                actual = str(tests.cell(row, 7).value or "")
                if not actual or "待实测" in actual:
                    errors.append(f"数据自测第{row}行仍待实测")

    all_text = "\n".join(
        str(cell.value or "")
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
    )
    if not allow_pending and "待实测" in all_text:
        errors.append("工作簿仍包含“待实测”")
    wb.close()
    return errors


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="基于标准模板生成或校验数仓模型交付文档。")
    parser.add_argument("--package", type=Path, help="JSON 工作包；生成模式必填")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--allow-pending", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        if not args.validate_only:
            if not args.package:
                raise ValueError("生成模式必须提供 --package。")
            package = json.loads(args.package.read_text(encoding="utf-8"))
            build_workbook(package, args.template, args.output, args.overwrite)
        if not args.output.exists():
            raise FileNotFoundError(f"待校验工作簿不存在：{args.output}")
        errors = validate_workbook(args.output, args.allow_pending)
    except (FileNotFoundError, FileExistsError, KeyError, ValueError, json.JSONDecodeError) as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 1
    if errors:
        for error in errors:
            print(f"[FAIL] {error}")
        return 1
    print(f"[PASS] {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
